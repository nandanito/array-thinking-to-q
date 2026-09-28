from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

wb = Workbook()
ws = wb.active
for row in [("Name", "Score"), ("Ada", 91), ("Grace", 88), ("Linus", 79)]:
    ws.append(row)
for cell in ws[1]:
    cell.font = Font(bold=True)
ws.freeze_panes = "A2"
wb.save("report.xlsx")

ws = load_workbook("report.xlsx").active
print(ws.freeze_panes)
print(ws["A1"].font.bold)
print(ws.max_row - 1)
