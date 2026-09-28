import csv, io
from openpyxl import Workbook, load_workbook

CSV = """region,product,qty
East,widget,4
West,gadget,7
East,gizmo,2
North,widget,5
West,widget,1
"""
wb = Workbook()
wb.remove(wb.active)
for row in csv.DictReader(io.StringIO(CSV)):
    if row["region"] not in wb.sheetnames:
        wb.create_sheet(row["region"]).append(["product", "qty"])
    wb[row["region"]].append([row["product"], int(row["qty"])])
wb.save("orders.xlsx")

for ws in load_workbook("orders.xlsx"):
    print(f"{ws.title}:{sum(r[1] for r in ws.iter_rows(min_row=2, values_only=True))}")
