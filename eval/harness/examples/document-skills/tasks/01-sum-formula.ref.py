from openpyxl import Workbook, load_workbook

wb = Workbook()
ws = wb.active
ws.title = "Sales"
ws.append(["Region", "Units"])
for row in [("North", 120), ("South", 95), ("East", 143), ("West", 88)]:
    ws.append(row)
ws["A6"], ws["B6"] = "Total", "=SUM(B2:B5)"
wb.save("sales.xlsx")

ws = load_workbook("sales.xlsx")["Sales"]
print(ws["B6"].value)
print(sum(ws.cell(row=r, column=2).value for r in range(2, 6)))
